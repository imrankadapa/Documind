import json
import re
import uuid
from pathlib import Path

import faiss
import fitz
import numpy as np
from docx import Document

from embedding_model import load_embedding_model

model = load_embedding_model()
BACKEND = Path(__file__).resolve().parents[1]
TEXT_FILE = BACKEND / "texts.npy"
INDEX_FILE = BACKEND / "vector.index"
INFO_FILE = BACKEND / "documents.json"


def split_text(text, size=800):
    text = re.sub(r"\s+", " ", text).strip()
    return [text[i:i + size] for i in range(0, len(text), size) if text[i:i + size].strip()]


def txt_extract(path):
    text = path.read_text(encoding="utf-8-sig").strip()
    return [{"text": text}] if text else []


def pdf_extract(path):
    pages = []
    pdf = fitz.open(str(path))
    for number, page in enumerate(pdf, 1):
        text = page.get_text().strip()
        if text:
            pages.append({"text": text, "page": number})
    pdf.close()
    return pages


def doc_extract(path):
    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return [{"text": "\n".join(parts)}]


def load_texts():
    if not TEXT_FILE.exists():
        return []
    return np.load(TEXT_FILE, allow_pickle=True).tolist()


def save_texts(texts):
    if not texts:
        return

    embeddings = model.encode([item["text"] for item in texts]).astype("float32")
    dimension = embeddings.shape[1]
    index = faiss.IndexFlatL2(dimension)
    index.add(embeddings)
    faiss.write_index(index, str(INDEX_FILE))
    np.save(TEXT_FILE, np.array(texts, dtype=object))


def save_document(pages, filename, file_type):
    document_id = uuid.uuid4().hex
    chunks = []
    chunk_number = 0

    for page in pages:
        for text in split_text(page["text"]):
            chunk_number += 1
            chunks.append({
                "text": text,
                "document_id": document_id,
                "filename": filename,
                "file_type": file_type,
                "chunk_number": chunk_number,
                **({"page": page["page"]} if "page" in page else {}),
            })

    if not chunks:
        raise ValueError("The document does not contain readable text.")

    save_texts(load_texts() + chunks)

    info = {
        "document_id": document_id,
        "filename": filename,
        "file_type": file_type,
        "chunks": len(chunks),
    }

    documents = json.loads(INFO_FILE.read_text(encoding="utf-8")) if INFO_FILE.exists() else []
    documents.append(info)
    INFO_FILE.write_text(json.dumps(documents, indent=2), encoding="utf-8")
    return info


def search_documents(question, document_ids=None, count=5):

    texts = load_texts()

    if not texts or not INDEX_FILE.exists():
        return []

    index = faiss.read_index(str(INDEX_FILE))
    question_vector = model.encode([question]).astype("float32")
    total = min(index.ntotal, len(texts))
    wanted = set(document_ids or [])
    candidate_count = total if wanted else min(count * 10, total)
    distances, positions = index.search(question_vector, candidate_count)
    results = []

    for distance, position in zip(distances[0], positions[0]):
        if position < 0 or position >= len(texts):
            continue
        item = dict(texts[position])
        if not wanted or item.get("document_id") in wanted:
            item["rank"] = len(results) + 1
            item["distance"] = round(float(distance), 4)
            location = item.get("filename", "document")
            if item.get("page"):
                location += f", page {item['page']}"
            location += f", chunk {item.get('chunk_number', '?')}"
            item["citation"] = f"[Source {item['rank']}] {location}"
            results.append(item)
        if len(results) == count:
            break
    return results
